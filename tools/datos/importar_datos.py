#!/usr/bin/env python3
"""Importa la planilla de datos del gimnasio (ver generar_planilla.py) al sistema Ikigai a traves de su API.

Pasos seguros, en este orden:
  1. VALIDAR   la planilla completa sin tocar nada; si hay errores, los lista (hoja y fila) y termina.
  2. PLANEAR   (con --url) compara con lo que ya existe y muestra que se crearia y que se omite.
  3. APLICAR   (con --aplicar) crea lo que falta. Se puede repetir: lo que ya existe se omite, no se duplica.

Ejemplos:
  python importar_datos.py Planilla.xlsx                                           # solo valida
  python importar_datos.py Planilla.xlsx --url https://app.ejemplo.com             # valida y muestra el plan
  python importar_datos.py Planilla.xlsx --url https://app.ejemplo.com --aplicar --crear-accesos --generar-clases 30

Usa la cuenta de administracion (usuario y clave: --usuario / variable IKIGAI_USUARIO y IKIGAI_CLAVE, o los pide).
La cuenta tiene que haber cambiado ya su clave temporal. Solo biblioteca estandar + openpyxl (requirements-datos.txt).
"""
import argparse
import datetime as dt
import getpass
import hashlib
import html
import http.cookiejar
import json
import os
import re
import sys
import time
import unicodedata
import urllib.error
import urllib.request
from decimal import Decimal, InvalidOperation

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from esquema import DIAS, ESTADOS, HOJAS, ORDEN, PERIODOS, SALDOS  # noqa: E402

HOY = dt.date.today()
EMAIL = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")


# ------------------------------------------------------------------ limpieza de valores
def norm(s):
    s = unicodedata.normalize("NFKD", str(s or "")).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", s).strip().casefold()


def texto(v):
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = str(v).strip()
    return s or None


def mayuscula_inicial(s):
    return s[:1].upper() + s[1:]


def solo_digitos(v):
    s = texto(v)
    return re.sub(r"\D", "", s) if s else None


def fecha(v):
    """Devuelve (date|None, ok). ok=False si habia algo escrito que no es una fecha."""
    if v is None or (isinstance(v, str) and not v.strip()):
        return None, True
    if isinstance(v, dt.datetime):
        return v.date(), True
    if isinstance(v, dt.date):
        return v, True
    if isinstance(v, (int, float)) and 20000 < v < 80000:                       # numero de serie de Excel
        return dt.date(1899, 12, 30) + dt.timedelta(days=int(v)), True
    s = str(v).strip()
    for f in ("%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d", "%d/%m/%y", "%d.%m.%Y"):
        try:
            return dt.datetime.strptime(s, f).date(), True
        except ValueError:
            pass
    return None, False


def hora(v):
    if v is None or (isinstance(v, str) and not v.strip()):
        return None, True
    if isinstance(v, dt.datetime):
        v = v.time()
    if isinstance(v, dt.time):
        return f"{v.hour:02d}:{v.minute:02d}", True
    if isinstance(v, float) and 0 <= v < 1:                                       # fraccion de dia de Excel
        minutos = round(v * 24 * 60)
        return f"{minutos // 60:02d}:{minutos % 60:02d}", True
    m = re.fullmatch(r"(\d{1,2})\s*[:.hH]\s*(\d{2})?\s*(?:hs?|HS?)?", str(v).strip())
    if m and int(m.group(1)) < 24 and int(m.group(2) or 0) < 60:
        return f"{int(m.group(1)):02d}:{int(m.group(2) or 0):02d}", True
    return None, False


def dinero(v):
    """Devuelve (Decimal|None, ok). Acepta 28000, '28.000', '28000,50', '$ 28.000'."""
    if v is None or (isinstance(v, str) and not v.strip()):
        return None, True
    if isinstance(v, (int, float, Decimal)):
        d = Decimal(str(v))
    else:
        s = re.sub(r"[^\d,.\-]", "", str(v))
        if "," in s:
            s = s.replace(".", "").replace(",", ".")
        elif s.count(".") > 1 or re.fullmatch(r"\d{1,3}(\.\d{3})+", s):
            s = s.replace(".", "")
        try:
            d = Decimal(s)
        except InvalidOperation:
            return None, False
    if d < 0:
        return None, False
    return d.quantize(Decimal("0.01")), True


def entero(v):
    if v is None or (isinstance(v, str) and not v.strip()):
        return None, True
    try:
        f = float(str(v).replace(",", "."))
    except ValueError:
        return None, False
    return (int(f), True) if f.is_integer() else (None, False)


def money_str(d):
    return str(d.quantize(Decimal("0.01")))


# ------------------------------------------------------------------ lectura de la planilla
def leer_planilla(ruta):
    from openpyxl import load_workbook          # importa aca para que --ayuda funcione sin openpyxl
    wb = load_workbook(ruta, data_only=True)
    datos, faltan = {}, []
    for nombre in ORDEN:
        if nombre not in wb.sheetnames:
            if nombre in ("Gimnasio", "Servicios", "Profesores", "Miembros", "Horarios"):
                faltan.append(f"Falta la hoja «{nombre}».")
            datos[nombre] = []
            continue
        ws = wb[nombre]
        encabezados = {}
        for j, celda in enumerate(ws[1], start=1):
            if celda.value is not None:
                encabezados[norm(str(celda.value).replace("*", ""))] = j
        cols = HOJAS[nombre]["columnas"]
        indice = {}
        for clave, titulo, obligatoria, *_ in cols:
            j = encabezados.get(norm(titulo.replace("*", "")))
            if j is None:
                faltan.append(f"En la hoja «{nombre}» falta la columna «{titulo}» (no cambies los encabezados).")
            indice[clave] = j
        filas = []
        for fila in range(2, ws.max_row + 1):
            registro = {clave: (ws.cell(row=fila, column=j).value if j else None) for clave, j in indice.items()}
            if all(texto(v) is None for v in registro.values()):
                continue
            registro["_fila"] = fila
            filas.append(registro)
        datos[nombre] = filas
    return datos, faltan


# ------------------------------------------------------------------ validacion
class Informe:
    def __init__(self):
        self.errores, self.avisos = [], []

    def error(self, hoja, fila, msg):
        self.errores.append(f"[{hoja}, fila {fila}] {msg}")

    def aviso(self, hoja, fila, msg):
        self.avisos.append(f"[{hoja}, fila {fila}] {msg}")


def validar(datos):
    """Devuelve (modelo, informe). El modelo es la planilla ya limpia y con las referencias resueltas."""
    inf = Informe()
    m = {k: [] for k in ORDEN}

    for r in datos["Gimnasio"][:1]:
        m["Gimnasio"].append({k: texto(r.get(k)) for k in ("nombre", "email", "telefono", "direccion", "horarios")})
        if m["Gimnasio"][0]["email"] and not EMAIL.match(m["Gimnasio"][0]["email"]):
            inf.error("Gimnasio", r["_fila"], "El correo no parece válido.")
    if len(datos["Gimnasio"]) > 1:
        inf.aviso("Gimnasio", datos["Gimnasio"][1]["_fila"], "Solo se usa la primera fila.")

    # ---- servicios
    servicios = {}
    for r in datos["Servicios"]:
        f, nombre = r["_fila"], texto(r.get("nombre"))
        precio, ok = dinero(r.get("precio"))
        if not nombre:
            inf.error("Servicios", f, "Falta el nombre del servicio.")
            continue
        if not ok:
            inf.error("Servicios", f, f"El precio «{r.get('precio')}» no es un importe válido (solo números).")
        if norm(nombre) in servicios:
            inf.error("Servicios", f, f"El servicio «{nombre}» está repetido.")
            continue
        s = {"nombre": nombre, "descripcion": texto(r.get("descripcion")), "precio": precio, "fila": f}
        servicios[norm(nombre)] = s
        m["Servicios"].append(s)

    # ---- profesores
    profesores = {}
    for r in datos["Profesores"]:
        f, nombre, disc = r["_fila"], texto(r.get("nombre")), texto(r.get("disciplina"))
        if not nombre or not disc:
            inf.error("Profesores", f, "Faltan el nombre y/o la disciplina.")
            continue
        email = texto(r.get("email"))
        if email and not EMAIL.match(email):
            inf.error("Profesores", f, f"El correo «{email}» no parece válido.")
        if norm(nombre) in profesores:
            inf.error("Profesores", f, f"El profesor «{nombre}» está repetido.")
            continue
        if norm(disc) not in servicios:
            inf.aviso("Profesores", f, f"La disciplina «{disc}» no figura en la hoja Servicios.")
        p = {"nombre": nombre, "disciplina": disc, "email": email, "telefono": texto(r.get("telefono")), "fila": f}
        profesores[norm(nombre)] = p
        m["Profesores"].append(p)

    # ---- miembros
    por_nombre, por_dni = {}, {}
    estados = {norm(k): v for k, v in ESTADOS.items()}
    for r in datos["Miembros"]:
        f, nombre = r["_fila"], texto(r.get("nombre"))
        if not nombre:
            inf.error("Miembros", f, "Falta el nombre completo.")
            continue
        dni = solo_digitos(r.get("dni"))
        if dni and not 6 <= len(dni) <= 9:
            inf.aviso("Miembros", f, f"El DNI «{r.get('dni')}» de {nombre} tiene una cantidad rara de dígitos.")
        email = texto(r.get("email"))
        if email and not EMAIL.match(email):
            inf.error("Miembros", f, f"El correo «{email}» de {nombre} no parece válido.")
        est_txt = texto(r.get("estado"))
        estado = estados.get(norm(est_txt)) if est_txt else "active"
        if estado is None:
            inf.error("Miembros", f, f"El estado «{est_txt}» no es válido (Activo, En prueba, Deudor, Suspendido o Baja).")
            estado = "active"
        alta, ok = fecha(r.get("alta"))
        if not ok:
            inf.error("Miembros", f, f"La fecha de alta «{r.get('alta')}» no es válida (dd/mm/aaaa).")
        if dni and dni in por_dni:
            inf.error("Miembros", f, f"El DNI {dni} ya está en la fila {por_dni[dni]['fila']} ({por_dni[dni]['nombre']}).")
            continue
        miembro = {"nombre": nombre, "dni": dni, "email": email, "telefono": texto(r.get("telefono")),
                   "direccion": texto(r.get("direccion")), "emergencia_nombre": texto(r.get("emergencia_nombre")),
                   "emergencia_telefono": texto(r.get("emergencia_telefono")), "estado": estado, "alta": alta,
                   "notas": texto(r.get("notas")), "fila": f}
        if dni:
            por_dni[dni] = miembro
        por_nombre.setdefault(norm(nombre), []).append(miembro)
        m["Miembros"].append(miembro)
    for nombre, lista in por_nombre.items():
        if len(lista) > 1 and any(x["dni"] is None for x in lista):
            for x in lista:
                inf.error("Miembros", x["fila"], f"Hay {len(lista)} personas llamadas «{x['nombre']}»: completá el DNI de todas para distinguirlas.")

    def miembro_de(valor, hoja, fila):
        t = texto(valor)
        if not t:
            inf.error(hoja, fila, "Falta indicar el miembro.")
            return None
        if solo_digitos(t) and len(solo_digitos(t)) >= 6 and solo_digitos(t) in por_dni:
            return por_dni[solo_digitos(t)]
        lista = por_nombre.get(norm(t), [])
        if len(lista) == 1:
            return lista[0]
        inf.error(hoja, fila, f"El miembro «{t}» " + ("es ambiguo (hay varios con ese nombre: usá el DNI)." if lista else "no figura en la hoja Miembros."))
        return None

    # ---- horarios
    dias = {norm(d): i for i, d in enumerate(DIAS)}
    horarios = {}
    for r in datos["Horarios"]:
        f = r["_fila"]
        serv, prof = texto(r.get("servicio")), texto(r.get("profesor"))
        dia = dias.get(norm(r.get("dia")))
        h, okh = hora(r.get("hora"))
        dur, okd = entero(r.get("duracion"))
        cupo, okc = entero(r.get("cupo"))
        problemas = []
        if not serv or norm(serv) not in servicios:
            problemas.append(f"el servicio «{serv or ''}» no está en la hoja Servicios")
        if not prof or norm(prof) not in profesores:
            problemas.append(f"el profesor «{prof or ''}» no está en la hoja Profesores")
        if dia is None:
            problemas.append(f"el día «{r.get('dia') or ''}» no es válido (Lunes a Domingo)")
        if h is None or not okh:
            problemas.append("la hora no es válida (formato 19:00)")
        if dur is None or not okd or not 1 <= dur <= 480:
            problemas.append("la duración tiene que ser un número de minutos entre 1 y 480")
        if cupo is None or not okc or not 1 <= cupo <= 30:
            problemas.append("el cupo tiene que ser un número entre 1 y 30")
        if problemas:
            inf.error("Horarios", f, mayuscula_inicial("; ".join(problemas)) + ".")
            continue
        clave = (norm(serv), norm(prof), dia, h)
        if clave in horarios:
            inf.error("Horarios", f, f"Este horario ya está en la fila {horarios[clave]['fila']}.")
            continue
        hh = {"servicio": servicios[norm(serv)], "profesor": profesores[norm(prof)], "dia": dia, "hora": h,
              "duracion": dur, "cupo": cupo, "fila": f, "clave": clave}
        horarios[clave] = hh
        m["Horarios"].append(hh)

    # ---- turnos fijos
    vistos = set()
    for r in datos["Turnos fijos"]:
        f = r["_fila"]
        mi = miembro_de(r.get("miembro"), "Turnos fijos", f)
        serv, prof = texto(r.get("servicio")), texto(r.get("profesor"))
        dia = dias.get(norm(r.get("dia")))
        h, okh = hora(r.get("hora"))
        desde, okf = fecha(r.get("desde"))
        if not okf:
            inf.error("Turnos fijos", f, f"La fecha «{r.get('desde')}» no es válida (dd/mm/aaaa).")
        clave = (norm(serv), norm(prof), dia, h)
        if mi is None:
            continue
        if clave not in horarios:
            inf.error("Turnos fijos", f, "No existe ese horario en la hoja Horarios (servicio, profesor, día y hora tienen que coincidir).")
            continue
        if (id(mi), clave) in vistos:
            inf.aviso("Turnos fijos", f, f"{mi['nombre']} está repetido en este horario.")
            continue
        vistos.add((id(mi), clave))
        m["Turnos fijos"].append({"miembro": mi, "horario": horarios[clave], "desde": desde or HOY, "fila": f})

    # ---- membresias
    periodos = {norm(k): v for k, v in PERIODOS.items()}
    vistas = set()
    for r in datos["Membresías"]:
        f = r["_fila"]
        mi = miembro_de(r.get("miembro"), "Membresías", f)
        serv = texto(r.get("servicio"))
        desde, okd = fecha(r.get("desde"))
        hasta, okh = fecha(r.get("hasta"))
        per_txt = texto(r.get("periodo"))
        periodo = periodos.get(norm(per_txt)) if per_txt else "monthly"
        precio, okp = dinero(r.get("precio"))
        problemas = []
        if not serv or norm(serv) not in servicios:
            problemas.append(f"el servicio «{serv or ''}» no está en la hoja Servicios")
        if desde is None or not okd:
            problemas.append("falta la fecha «Desde» o no es válida (dd/mm/aaaa)")
        if not okh:
            problemas.append("la fecha «Hasta» no es válida")
        elif desde and hasta and hasta < desde:
            problemas.append("«Hasta» es anterior a «Desde»")
        if periodo is None:
            problemas.append(f"el período «{per_txt}» no es válido")
        if not okp:
            problemas.append("el precio no es un importe válido")
        if serv and norm(serv) in servicios and precio is None and servicios[norm(serv)]["precio"] is None:
            problemas.append("falta el precio (ni la membresía ni el servicio lo tienen)")
        if mi is None or problemas:
            if problemas:
                inf.error("Membresías", f, mayuscula_inicial("; ".join(problemas)) + ".")
            continue
        s = servicios[norm(serv)]
        if (id(mi), id(s), desde) in vistas:
            inf.aviso("Membresías", f, f"{mi['nombre']} ya tiene una membresía de {s['nombre']} que empieza ese día; se ignora la repetida.")
            continue
        vistas.add((id(mi), id(s), desde))
        m["Membresías"].append({"miembro": mi, "servicio": s, "desde": desde, "hasta": hasta, "periodo": periodo,
                                "precio": precio if precio is not None else s["precio"], "fila": f})

    # ---- saldos iniciales
    tipos = {norm(k): v for k, v in SALDOS.items()}
    for r in datos["Saldos iniciales"]:
        f = r["_fila"]
        mi = miembro_de(r.get("miembro"), "Saldos iniciales", f)
        tipo = tipos.get(norm(r.get("tipo")))
        imp, oki = dinero(r.get("importe"))
        fch, okf = fecha(r.get("fecha"))
        problemas = []
        if tipo is None:
            problemas.append("el tipo tiene que ser «Debe» o «A favor»")
        if imp is None or not oki or imp <= 0:
            problemas.append("el importe tiene que ser un número mayor que cero")
        if fch is None or not okf:
            problemas.append("falta la fecha o no es válida (dd/mm/aaaa)")
        if mi is None or problemas:
            if problemas:
                inf.error("Saldos iniciales", f, mayuscula_inicial("; ".join(problemas)) + ".")
            continue
        m["Saldos iniciales"].append({"miembro": mi, "tipo": tipo, "importe": imp, "fecha": fch,
                                      "detalle": texto(r.get("detalle")) or "Saldo inicial al empezar a usar el sistema", "fila": f})

    if not (m["Servicios"] or m["Profesores"] or m["Miembros"]):
        inf.errores.append("La planilla está vacía: completá al menos las hojas Servicios, Profesores y Miembros.")

    # ---- avisos de consistencia
    sin_membresia = [x["nombre"] for x in m["Miembros"] if x["estado"] in ("active", "debtor") and not any(ms["miembro"] is x for ms in m["Membresías"])]
    if m["Membresías"] and sin_membresia:
        inf.avisos.append(f"[Membresías] {len(sin_membresia)} miembro(s) activos o deudores no tienen membresía: {', '.join(sin_membresia[:5])}{'…' if len(sin_membresia) > 5 else ''}.")
    return m, inf


# ------------------------------------------------------------------ cliente de la API
class ApiError(Exception):
    def __init__(self, status, cuerpo, ruta):
        super().__init__(f"HTTP {status} en {ruta}: {cuerpo}")
        self.status, self.cuerpo = status, cuerpo


class Api:
    def __init__(self, origen, api=None):
        self.origen = origen.rstrip("/")
        self.base = (api or self.origen + "/api/v1").rstrip("/")
        self.csrf = None
        self.opener = urllib.request.build_opener(urllib.request.HTTPCookieProcessor(http.cookiejar.CookieJar()))

    def llamar(self, metodo, ruta, cuerpo=None, encabezados=None, reintentos=6):
        cab = {"Origin": self.origen, "Content-Type": "application/json", "Accept": "application/json"}
        if self.csrf and metodo != "GET":
            cab["X-CSRF-Token"] = self.csrf
        cab.update(encabezados or {})
        datos = json.dumps(cuerpo).encode() if cuerpo is not None else None
        for intento in range(reintentos):
            req = urllib.request.Request(self.base + ruta, data=datos, method=metodo, headers=cab)
            try:
                with self.opener.open(req, timeout=60) as resp:
                    texto_resp = resp.read()
                    return resp.status, (json.loads(texto_resp) if texto_resp else None)
            except urllib.error.HTTPError as e:
                crudo = e.read()
                try:
                    cuerpo_err = json.loads(crudo)
                except ValueError:
                    cuerpo_err = crudo.decode("utf-8", "replace")[:300]
                if e.code == 429 and intento < reintentos - 1:                      # limite de intentos: esperar y seguir
                    espera = min(int(e.headers.get("Retry-After", "20")), 90)
                    print(f"    … el servidor pidió esperar {espera}s (límite de velocidad)", flush=True)
                    time.sleep(espera + 1)
                    continue
                return e.code, cuerpo_err
            except urllib.error.URLError as e:
                if intento < 2:
                    time.sleep(2)
                    continue
                raise SystemExit(f"No se pudo conectar con {self.base}: {e.reason}")
        raise SystemExit("Demasiados reintentos.")

    def ok(self, metodo, ruta, cuerpo=None, encabezados=None, aceptar=(200, 201)):
        estado, resp = self.llamar(metodo, ruta, cuerpo, encabezados)
        if estado not in aceptar:
            raise ApiError(estado, resp, f"{metodo} {ruta}")
        return resp

    def entrar(self, usuario, clave):
        estado, resp = self.llamar("POST", "/auth/login", {"login": usuario, "password": clave})
        if estado != 200:
            raise SystemExit(f"No se pudo iniciar sesión ({estado}): {resp}")
        if resp.get("mustChangePassword"):
            raise SystemExit("Esa cuenta todavía tiene la clave temporal: entrá una vez al portal, elegí tu clave y volvé a correr el script.")
        if "admin" not in resp["user"]["roles"]:
            raise SystemExit("La cuenta no es de administración.")
        self.csrf = resp["csrfToken"]


# ------------------------------------------------------------------ aplicar
class Importador:
    def __init__(self, api, modelo, aplicar):
        self.api, self.m, self.aplicar = api, modelo, aplicar
        self.cuenta = {}                      # resumen: paso -> [creados, omitidos]
        self.ids = {"serv": {}, "prof": {}, "miem": {}, "hor": {}}
        self.credenciales = []

    def _c(self, paso, creado):
        c = self.cuenta.setdefault(paso, [0, 0])
        c[0 if creado else 1] += 1

    def run(self, crear_accesos=False, generar_dias=0, archivo_credenciales=None, url_app=""):
        self.gimnasio()
        self.servicios()
        self.profesores()
        self.miembros()
        self.horarios()
        self.turnos()
        self.membresias()
        self.saldos()
        if generar_dias:
            self.generar_clases(generar_dias)
        if crear_accesos:
            self.accesos(archivo_credenciales, url_app)

    def gimnasio(self):
        g = self.m["Gimnasio"][0] if self.m["Gimnasio"] else None
        if not g or not any(g.values()):
            return
        cuerpo = {k: v for k, v in {"institution_name": g["nombre"], "contact_email": g["email"], "contact_phone": g["telefono"],
                                    "address": g["direccion"], "opening_hours": g["horarios"]}.items() if v}
        if cuerpo and self.aplicar:
            self.api.ok("PATCH", "/admin/institution-settings", cuerpo)
        self._c("Datos del gimnasio", True)

    def servicios(self):
        existentes = {norm(s["name"]): s for s in self.api.ok("GET", "/services")["services"]} if self.api else {}
        for s in self.m["Servicios"]:
            ya = existentes.get(norm(s["nombre"]))
            if ya:
                self.ids["serv"][norm(s["nombre"])] = ya["id"]
                self._c("Servicios", False)
                continue
            if self.aplicar:
                cuerpo = {"name": s["nombre"]}
                if s["descripcion"]:
                    cuerpo["description"] = s["descripcion"]
                if s["precio"] is not None:
                    cuerpo["currentPrice"] = money_str(s["precio"])
                self.ids["serv"][norm(s["nombre"])] = self.api.ok("POST", "/services", cuerpo)["service"]["id"]
            else:
                self.ids["serv"][norm(s["nombre"])] = "(nuevo)"
            self._c("Servicios", True)

    def profesores(self):
        existentes = {norm(c["full_name"]): c for c in self.api.ok("GET", "/coaches")["coaches"]} if self.api else {}
        for p in self.m["Profesores"]:
            ya = existentes.get(norm(p["nombre"]))
            if ya:
                self.ids["prof"][norm(p["nombre"])] = ya["id"]
                self._c("Profesores", False)
                continue
            if self.aplicar:
                cuerpo = {"full_name": p["nombre"], "discipline": p["disciplina"]}
                if p["email"]:
                    cuerpo["email"] = p["email"]
                if p["telefono"]:
                    cuerpo["phone"] = p["telefono"]
                self.ids["prof"][norm(p["nombre"])] = self.api.ok("POST", "/coaches", cuerpo)["coach"]["id"]
            else:
                self.ids["prof"][norm(p["nombre"])] = "(nuevo)"
            self._c("Profesores", True)

    def _miembros_existentes(self):
        todos, offset = [], 0
        while self.api:
            pagina = self.api.ok("GET", f"/members?limit=100&offset={offset}")
            todos += pagina["items"]
            offset += 100
            if offset >= pagina["total"]:
                break
        return todos

    def miembros(self):
        existentes = self._miembros_existentes()
        por_dni = {solo_digitos(x["dni"]): x for x in existentes if x.get("dni")}
        por_nombre = {}
        for x in existentes:
            por_nombre.setdefault(norm(x["full_name"]), []).append(x)
        for mi in self.m["Miembros"]:
            ya = por_dni.get(mi["dni"]) if mi["dni"] else None
            if not ya:
                candidatos = [x for x in por_nombre.get(norm(mi["nombre"]), []) if not (mi["dni"] and x.get("dni"))]
                ya = candidatos[0] if len(candidatos) == 1 else None
            if ya:
                self.ids["miem"][id(mi)] = ya["id"]
                self._c("Miembros", False)
                continue
            if self.aplicar:
                cuerpo = {"full_name": mi["nombre"], "status": mi["estado"]}
                for k_api, k in (("email", "email"), ("phone", "telefono"), ("address", "direccion"), ("dni", "dni"),
                                 ("emergency_contact_name", "emergencia_nombre"), ("emergency_contact_phone", "emergencia_telefono"),
                                 ("notes", "notas")):
                    if mi[k]:
                        cuerpo[k_api] = mi[k]
                if mi["alta"]:
                    cuerpo["joined_on"] = mi["alta"].isoformat()
                if mi["estado"] == "exited":
                    cuerpo["exited_on"] = HOY.isoformat()
                self.ids["miem"][id(mi)] = self.api.ok("POST", "/members", cuerpo)["member"]["id"]
            else:
                self.ids["miem"][id(mi)] = "(nuevo)"
            self._c("Miembros", True)

    def horarios(self):
        existentes = {}
        if self.api:
            for t in self.api.ok("GET", "/class-templates")["templates"]:
                existentes[(norm(t["service"]["name"]), norm(t["coach"]["name"]), t["weekday"], t["localStartTime"][:5])] = t["id"]
        for h in self.m["Horarios"]:
            clave = h["clave"]
            if clave in existentes:
                self.ids["hor"][clave] = existentes[clave]
                self._c("Horarios", False)
                continue
            if self.aplicar:
                cuerpo = {"service_id": self.ids["serv"][clave[0]], "coach_id": self.ids["prof"][clave[1]], "weekday": h["dia"],
                          "local_start_time": h["hora"], "duration_minutes": h["duracion"], "capacity": h["cupo"]}
                self.ids["hor"][clave] = self.api.ok("POST", "/class-templates", cuerpo)["template"]["id"]
            else:
                self.ids["hor"][clave] = "(nuevo)"
            self._c("Horarios", True)

    def turnos(self):
        cache = {}
        for t in self.m["Turnos fijos"]:
            hid = self.ids["hor"][t["horario"]["clave"]]
            mid = self.ids["miem"][id(t["miembro"])]
            if hid == "(nuevo)" or mid == "(nuevo)":
                self._c("Turnos fijos", True)
                continue
            if hid not in cache:
                cache[hid] = {a["memberId"] for a in self.api.ok("GET", f"/class-templates/{hid}/assignments")["assignments"] if a.get("active", True)}
            if mid in cache[hid]:
                self._c("Turnos fijos", False)
                continue
            if self.aplicar:
                self.api.ok("POST", f"/class-templates/{hid}/assignments", {"member_id": mid, "starts_on": t["desde"].isoformat()})
                cache[hid].add(mid)
            self._c("Turnos fijos", True)

    def membresias(self):
        for ms in self.m["Membresías"]:
            mid, sid = self.ids["miem"][id(ms["miembro"])], self.ids["serv"][norm(ms["servicio"]["nombre"])]
            if mid == "(nuevo)" or sid == "(nuevo)":
                self._c("Membresías", True)
                continue
            lista = self.api.ok("GET", f"/members/{mid}/memberships")["memberships"]
            if any(x["service_id"] == sid and x["starts_on"][:10] == ms["desde"].isoformat() for x in lista):
                self._c("Membresías", False)
                continue
            if self.aplicar:
                cuerpo = {"serviceId": sid, "startsOn": ms["desde"].isoformat(), "billingPeriod": ms["periodo"], "agreedPrice": money_str(ms["precio"])}
                if ms["hasta"]:
                    cuerpo["endsOn"] = ms["hasta"].isoformat()
                self.api.ok("POST", f"/members/{mid}/memberships", cuerpo)
            self._c("Membresías", True)

    def saldos(self):
        for s in self.m["Saldos iniciales"]:
            mid = self.ids["miem"][id(s["miembro"])]
            if mid == "(nuevo)" or not self.aplicar:
                self._c("Saldos iniciales", True)
                continue
            # La clave de idempotencia sale del contenido: repetir la importacion no duplica el saldo.
            clave = "import-" + hashlib.sha1(f"{mid}|{s['tipo']}|{s['importe']}|{s['fecha']}|{s['detalle']}".encode()).hexdigest()[:32]
            cuerpo = {"type": "adjustment", "direction": s["tipo"], "amount": money_str(s["importe"]), "description": s["detalle"][:240],
                      "occurred_on": s["fecha"].isoformat()}
            estado, resp = self.api.llamar("POST", f"/members/{mid}/ledger-entries", cuerpo, {"Idempotency-Key": clave})
            if estado not in (200, 201):
                raise ApiError(estado, resp, f"POST /members/{mid}/ledger-entries")
            self._c("Saldos iniciales", estado == 201)

    def generar_clases(self, dias):
        if not self.aplicar:
            self._c("Clases generadas (días)", True)
            return
        desde, hasta = HOY, HOY + dt.timedelta(days=dias - 1)
        r = self.api.ok("POST", "/class-sessions/generate", {"from_date": desde.isoformat(), "to_date": hasta.isoformat()})
        print(f"    clases creadas: {r['created']}, ya existían: {r['alreadyExisting']}")
        self._c("Clases generadas (días)", True)

    def accesos(self, archivo, url_app):
        if not self.aplicar:
            self._c("Accesos a crear", True)
            return
        cuentas = self.api.ok("GET", "/admin/accounts")["accounts"]
        mios_m = {self.ids["miem"][id(x)] for x in self.m["Miembros"] if x["estado"] in ("active", "trial", "debtor")}
        mios_p = set(self.ids["prof"].values())
        for c in cuentas:
            if c.get("user_id"):
                continue
            if c["role"] == "member" and c["profile_id"] not in mios_m:
                continue
            if c["role"] == "coach" and c["profile_id"] not in mios_p:
                continue
            cuerpo = {"role": c["role"], ("memberId" if c["role"] == "member" else "coachId"): c["profile_id"]}
            r = self.api.ok("POST", "/admin/accounts", cuerpo)
            registro = {"nombre": r["account"]["fullName"], "rol": c["role"], "usuario": r["account"]["username"],
                        "clave": r["temporaryPassword"]}
            self.credenciales.append(registro)
            with open(archivo, "a", encoding="utf-8") as f:                      # a disco al instante: la clave solo se ve una vez
                f.write(json.dumps(registro, ensure_ascii=False) + "\n")
            self._c("Accesos creados", True)
        self.cuenta.setdefault("Accesos creados", [0, 0])


# ------------------------------------------------------------------ hoja imprimible de accesos
def hoja_accesos(registros, url_app, destino):
    base = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

    def b64(ruta):
        import base64
        with open(os.path.join(base, ruta), "rb") as f:
            return base64.b64encode(f.read()).decode()
    logo = b64("assets/pwa/icon-192.png")
    fuentes = "".join(
        f"@font-face{{font-family:'{fam}';font-weight:{w};src:url(data:font/woff2;base64,{b64('assets/fonts/' + arch)}) format('woff2')}}"
        for fam, w, arch in (("Bebas Neue", 400, "bebas-neue-400.woff2"), ("Rajdhani", 600, "rajdhani-600.woff2"), ("Rajdhani", 700, "rajdhani-700.woff2")))
    rol = {"member": "Miembro", "coach": "Profesor"}
    registros = sorted(registros, key=lambda r: (r["rol"] != "coach", norm(r["nombre"])))   # profesores primero, luego A-Z
    tarjetas = "".join(f"""
    <article class="tarjeta">
      <header><img src="data:image/png;base64,{logo}" alt=""><div><strong>IKIGAI</strong><small>TRAINING</small></div></header>
      <p class="quien">{html.escape(r['nombre'])}<span>{rol.get(r['rol'], r['rol'])}</span></p>
      <dl><dt>Usuario</dt><dd>{html.escape(r['usuario'])}</dd><dt>Clave temporal</dt><dd>{html.escape(r['clave'])}</dd></dl>
      <ol><li>Abrí la app de Ikigai{(' o entrá a <b>' + html.escape(url_app) + '</b>') if url_app else ''}.</li>
      <li>Ingresá con el usuario y la clave de arriba.</li><li>Elegí tu propia clave (mínimo 12 caracteres). La temporal deja de servir.</li></ol>
    </article>""" for r in registros)
    contenido = f"""<!doctype html><html lang="es-AR"><head><meta charset="utf-8"><title>Accesos a Ikigai — {HOY:%d/%m/%Y}</title>
<style>{fuentes}
@page{{size:A4;margin:10mm}} *{{box-sizing:border-box}} body{{margin:0;font-family:'Rajdhani',sans-serif;color:#111;font-size:12pt}}
.aviso{{background:#fff3cd;border:2px solid #e6b755;padding:8px 12px;margin:0 0 8mm;font-weight:700}}
.grilla{{display:grid;grid-template-columns:1fr 1fr;gap:6mm}}
.tarjeta{{border:1.5px dashed #888;border-radius:6px;padding:5mm;break-inside:avoid;min-height:62mm}}
header{{display:flex;align-items:center;gap:3mm;border-bottom:2px solid #e31a20;padding-bottom:2mm}} header img{{width:11mm;height:11mm;border-radius:2mm}}
header strong{{display:block;font:400 17pt 'Bebas Neue';letter-spacing:1px;line-height:1}} header small{{letter-spacing:3px;color:#666;font-size:8pt}}
.quien{{font-weight:700;font-size:13pt;margin:3mm 0 1mm}} .quien span{{display:block;color:#e31a20;font-size:9pt;text-transform:uppercase;letter-spacing:1.5px}}
dl{{margin:0}} dt{{font-size:8pt;color:#666;text-transform:uppercase;letter-spacing:1px;margin-top:1.5mm}}
dd{{margin:0;font:700 13pt ui-monospace,Menlo,Consolas,monospace;background:#f2f2f4;padding:1mm 2mm;border-radius:1mm}}
ol{{margin:2.5mm 0 0;padding-left:5mm;font-size:9pt;color:#333}}
@media print{{.aviso{{display:none}}}}
</style></head><body>
<p class="aviso">CONFIDENCIAL: este archivo contiene claves. Imprimilo, entregá cada tarjeta EN MANO a su dueño y después borrá el archivo (y la papelera).</p>
<main class="grilla">{tarjetas}</main></body></html>"""
    os.makedirs(os.path.dirname(destino) or ".", exist_ok=True)
    fd = os.open(destino, os.O_WRONLY | os.O_CREAT | os.O_TRUNC, 0o600)
    with os.fdopen(fd, "w", encoding="utf-8") as f:
        f.write(contenido)
    return destino


# ------------------------------------------------------------------ principal
def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("planilla", nargs="?", help="Archivo .xlsx completado.")
    ap.add_argument("--url", help="Direccion del sistema (la del portal), por ejemplo https://app.ejemplo.com")
    ap.add_argument("--api", help="Direccion de la API si no es <url>/api/v1 (por ejemplo en desarrollo).")
    ap.add_argument("--usuario", default=os.environ.get("IKIGAI_USUARIO"), help="Correo de la cuenta de administracion.")
    ap.add_argument("--aplicar", action="store_true", help="Ejecuta la carga (sin esto solo valida y muestra el plan).")
    ap.add_argument("--crear-accesos", action="store_true", help="Crea usuario y clave temporal a profesores y miembros que no tengan.")
    ap.add_argument("--generar-clases", type=int, default=0, metavar="DIAS", help="Genera las clases de los proximos DIAS dias.")
    ap.add_argument("--salida", default=os.path.expanduser("~/Ikigai-accesos"), help="Carpeta donde queda la hoja de accesos.")
    ap.add_argument("--hoja-desde", metavar="ARCHIVO.jsonl", help="Regenera la hoja imprimible de accesos desde un archivo guardado.")
    ap.add_argument("--si", action="store_true", help="No pedir confirmacion antes de aplicar.")
    a = ap.parse_args()

    if a.hoja_desde:
        regs = [json.loads(linea) for linea in open(a.hoja_desde, encoding="utf-8") if linea.strip()]
        print("Hoja escrita en", hoja_accesos(regs, a.url or "", a.hoja_desde.replace(".jsonl", ".html")))
        return
    if not a.planilla:
        ap.error("Indicá la planilla .xlsx")

    datos, faltan = leer_planilla(a.planilla)
    if faltan:
        print("La planilla no tiene la estructura esperada:")
        for f in faltan:
            print("  ✗", f)
        sys.exit(2)
    modelo, inf = validar(datos)
    print("== Validación de la planilla")
    for t in ORDEN:
        print(f"   {t:<18} {len(modelo[t]):>4} filas válidas")
    for av in inf.avisos:
        print("  ! aviso:", av)
    if inf.errores:
        print(f"\n{len(inf.errores)} ERROR(ES): corregí la planilla y volvé a correr. No se tocó nada.")
        for e in inf.errores:
            print("  ✗", e)
        sys.exit(1)
    print("   Sin errores.")
    if not a.url:
        print("\nPara ver el plan contra el sistema, agregá --url.")
        return

    api = Api(a.url, a.api)
    usuario = a.usuario or input("Correo de administración: ").strip()
    clave = os.environ.get("IKIGAI_CLAVE") or getpass.getpass("Clave: ")
    api.entrar(usuario, clave)
    print(f"\n== {'APLICANDO' if a.aplicar else 'Plan (no se escribe nada)'} en {a.url}")
    if a.aplicar and not a.si:
        if input("Se va a cargar la información en el sistema real. Escribí SI para continuar: ").strip().upper() != "SI":
            print("Cancelado.")
            return
    os.makedirs(a.salida, exist_ok=True)
    os.chmod(a.salida, 0o700)
    marca = dt.datetime.now().strftime("%Y%m%d-%H%M")
    jsonl = os.path.join(a.salida, f"accesos-{marca}.jsonl")
    imp = Importador(api, modelo, a.aplicar)
    try:
        imp.run(crear_accesos=a.crear_accesos, generar_dias=a.generar_clases, archivo_credenciales=jsonl, url_app=a.url)
    except ApiError as e:
        print(f"\n✗ El sistema rechazó una operación: {e}\n  Lo ya cargado queda (se puede volver a correr: no duplica).")
        sys.exit(3)
    finally:
        if imp.credenciales:
            os.chmod(jsonl, 0o600)
            destino = hoja_accesos(imp.credenciales, a.url, os.path.join(a.salida, f"accesos-{marca}.html"))
            print(f"\n== Accesos creados: {len(imp.credenciales)}\n   Hoja imprimible: {destino}\n   (contiene claves: imprimir, entregar en mano y borrar)")
    verbo = ("creados", "ya existían") if a.aplicar else ("a crear", "ya existen")
    print(f"\n== Resultado ({verbo[0]} / {verbo[1]})")
    for paso, (c, o) in imp.cuenta.items():
        print(f"   {paso:<26} {c:>4} / {o:>4}")
    if not a.aplicar:
        print("\nPara cargar de verdad, repetí el comando con --aplicar.")


if __name__ == "__main__":
    main()
