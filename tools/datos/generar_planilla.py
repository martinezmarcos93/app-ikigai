"""Genera la planilla Excel que se le pide completar al gimnasio (y la misma con datos ficticios, de ejemplo).

    python generar_planilla.py                   # plantillas/Planilla-Ikigai-VACIA.xlsx
    python generar_planilla.py --ejemplo         # plantillas/Planilla-Ikigai-EJEMPLO.xlsx (gimnasio ficticio)

Las columnas Servicio, Profesor, Miembro, Dia, Estado, Periodo y Tipo traen lista desplegable para evitar errores de
tipeo; las listas de Servicio/Profesor/Miembro salen de las hojas correspondientes.
"""
import argparse
import datetime as dt
import os
import sys

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from esquema import DIAS, ESTADOS, HOJAS, ORDEN, PERIODOS, SALDOS  # noqa: E402

ROJO = "E31A20"
NEGRO = "111112"
GRIS = "F2F2F4"
BORDE = Side(style="thin", color="C9C9CF")

LEEME = [
    ("PLANILLA DE DATOS DEL GIMNASIO", "titulo"),
    ("Esta planilla sirve para cargar de una sola vez los datos del gimnasio en la aplicación. Completala con calma: lo que anotes acá es lo que va a ver cada profesor y cada miembro.", None),
    ("", None),
    ("CÓMO COMPLETARLA", "sub"),
    ("1. Completá las hojas de izquierda a derecha: Gimnasio, Servicios, Profesores, Miembros, Horarios. Las demás (Turnos fijos, Membresías, Saldos iniciales) son opcionales pero muy recomendables.", None),
    ("2. Las columnas con asterisco (*) son obligatorias. El resto puede quedar vacío.", None),
    ("3. Donde hay una flechita al lado de la celda (Servicio, Profesor, Miembro, Día, Estado…), elegí el valor de la lista en vez de escribirlo: así no hay errores de tipeo.", None),
    ("4. Fechas como dd/mm/aaaa (por ejemplo 15/03/2026). Horas como 19:00. Importes solo con números, sin $ ni puntos (28000).", None),
    ("5. Una fila por persona o por cosa. No dejes filas vacías en el medio ni cambies el nombre de las hojas ni de las columnas.", None),
    ("6. Si dos personas tienen exactamente el mismo nombre, completá el DNI de ambas para poder distinguirlas.", None),
    ("", None),
    ("QUÉ HACE CADA HOJA", "sub"),
    ("Gimnasio: nombre, contacto y horarios de atención (una sola fila).", None),
    ("Servicios: las disciplinas que se ofrecen y su precio de lista.", None),
    ("Profesores: quién da clases y qué enseña.", None),
    ("Miembros: todas las personas que entrenan, con sus datos de contacto y estado (Activo, En prueba, Deudor, Suspendido, Baja).", None),
    ("Horarios: la grilla semanal de clases (qué servicio, con qué profesor, qué día y a qué hora, con qué cupo).", None),
    ("Turnos fijos: quién tiene su lugar reservado siempre en el mismo horario.", None),
    ("Membresías: qué servicio contrató cada miembro, desde cuándo y a qué precio.", None),
    ("Saldos iniciales: lo que cada persona debe o tiene a favor el día que se empieza a usar la aplicación.", None),
    ("", None),
    ("MUY IMPORTANTE: PRIVACIDAD", "sub"),
    ("Esta planilla contiene datos personales (nombres, DNI, teléfonos). Es solo para el uso del gimnasio: no la reenvíes por canales abiertos, guardala en un lugar seguro y borrala (también de la papelera y del correo) una vez cargados los datos. Anotá únicamente lo necesario; en «Notas internas» no escribas datos de salud ni opiniones sobre otras personas.", None),
    ("", None),
    ("Cuando termines, enviá el archivo por un medio de confianza. Se revisa, se carga y las claves de acceso se entregan en mano.", None),
]


def _estilo_encabezado(celda, obligatoria):
    celda.font = Font(bold=True, color="FFFFFF", size=11)
    celda.fill = PatternFill("solid", fgColor=ROJO if obligatoria else NEGRO)
    celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    celda.border = Border(top=BORDE, bottom=BORDE, left=BORDE, right=BORDE)


def _fecha(v):
    return v if isinstance(v, (dt.date, dt.datetime)) else v


def construir(datos=None, ejemplo=False):
    wb = Workbook()
    leeme = wb.active
    leeme.title = "Léeme"
    leeme.sheet_view.showGridLines = False
    leeme.column_dimensions["A"].width = 120
    for i, (texto, tipo) in enumerate(LEEME, start=1):
        c = leeme.cell(row=i, column=1, value=texto)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if tipo == "titulo":
            c.font = Font(bold=True, size=16, color=ROJO)
        elif tipo == "sub":
            c.font = Font(bold=True, size=12, color=NEGRO)
        if ejemplo and i == 2:
            c.value = "ESTA ES UNA PLANILLA DE EJEMPLO con un gimnasio inventado: sirve para ver cómo se completa. " + c.value
            c.font = Font(bold=True, color=ROJO)
        if len(texto) > 118:
            leeme.row_dimensions[i].height = 15 * (len(texto) // 118 + 1)

    hojas = {}
    for nombre in ORDEN:
        meta = HOJAS[nombre]
        ws = wb.create_sheet(nombre)
        hojas[nombre] = ws
        cols = meta["columnas"]
        for j, (clave, titulo, obligatoria, ancho, ayuda) in enumerate(cols, start=1):
            c = ws.cell(row=1, column=j, value=titulo)
            _estilo_encabezado(c, obligatoria)
            c.comment = Comment(ayuda, "Ikigai")
            ws.column_dimensions[get_column_letter(j)].width = ancho
        ws.row_dimensions[1].height = 32
        ws.freeze_panes = "A2"
        ws.sheet_properties.tabColor = ROJO if nombre in ("Gimnasio", "Servicios", "Profesores", "Miembros", "Horarios") else "888888"
        for fila in range(2, meta["maximo"] + 2):
            for j, (clave, *_rest) in enumerate(cols, start=1):
                ws.cell(row=fila, column=j).border = Border(bottom=BORDE)
                if clave in ("dni", "telefono", "emergencia_telefono"):
                    ws.cell(row=fila, column=j).number_format = "@"          # texto: no pierde ceros ni se vuelve cientifico
                elif clave in ("alta", "desde", "hasta", "fecha"):
                    ws.cell(row=fila, column=j).number_format = "dd/mm/yyyy"
                elif clave == "hora":
                    ws.cell(row=fila, column=j).number_format = "hh:mm"
                elif clave in ("precio", "importe"):
                    ws.cell(row=fila, column=j).number_format = "#,##0"

    def lista(hoja, clave, formula_o_valores, advertencia=False):
        cols = [k for k, *_ in HOJAS[hoja]["columnas"]]
        letra = get_column_letter(cols.index(clave) + 1)
        maximo = HOJAS[hoja]["maximo"] + 1
        dv = DataValidation(type="list", formula1=formula_o_valores, allow_blank=True,
                            errorStyle="warning" if advertencia else "stop",
                            errorTitle="Valor fuera de la lista", error="Elegí un valor de la lista desplegable.",
                            showErrorMessage=True)
        hojas[hoja].add_data_validation(dv)
        dv.add(f"{letra}2:{letra}{maximo}")

    dias = '"' + ",".join(DIAS[1:] + DIAS[:1]) + '"'
    nombres = {"servicio": "=Servicios!$A$2:$A$61", "profesor": "=Profesores!$A$2:$A$61", "miembro": "=Miembros!$A$2:$A$501"}
    for hoja in ("Horarios", "Turnos fijos", "Membresías"):
        lista(hoja, "servicio", nombres["servicio"], advertencia=True)
    for hoja in ("Horarios", "Turnos fijos"):
        lista(hoja, "profesor", nombres["profesor"], advertencia=True)
        lista(hoja, "dia", dias)
    for hoja in ("Turnos fijos", "Membresías", "Saldos iniciales"):
        lista(hoja, "miembro", nombres["miembro"], advertencia=True)
    lista("Miembros", "estado", '"' + ",".join(ESTADOS) + '"')
    lista("Membresías", "periodo", '"' + ",".join(PERIODOS) + '"')
    lista("Saldos iniciales", "tipo", '"' + ",".join(SALDOS) + '"')
    lista("Profesores", "disciplina", "=Servicios!$A$2:$A$61", advertencia=True)

    if datos:
        for nombre, filas in datos.items():
            ws = hojas[nombre]
            cols = [k for k, *_ in HOJAS[nombre]["columnas"]]
            for i, fila in enumerate(filas, start=2):
                for j, clave in enumerate(cols, start=1):
                    valor = fila.get(clave)
                    if clave == "hora" and isinstance(valor, str) and ":" in valor:
                        h, m = valor.split(":")[:2]
                        valor = dt.time(int(h), int(m))
                    ws.cell(row=i, column=j, value=_fecha(valor))
    return wb


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--ejemplo", action="store_true", help="Incluye el gimnasio ficticio de ejemplo.")
    ap.add_argument("--salida", help="Archivo de salida (por defecto en plantillas/).")
    a = ap.parse_args()
    base = os.path.dirname(os.path.abspath(__file__))
    datos = None
    if a.ejemplo:
        from ejemplo import DATOS
        datos = DATOS
    salida = a.salida or os.path.join(base, "plantillas", "Planilla-Ikigai-EJEMPLO.xlsx" if a.ejemplo else "Planilla-Ikigai-VACIA.xlsx")
    os.makedirs(os.path.dirname(salida), exist_ok=True)
    construir(datos, ejemplo=a.ejemplo).save(salida)
    print(f"Planilla escrita en {salida}")


if __name__ == "__main__":
    main()
